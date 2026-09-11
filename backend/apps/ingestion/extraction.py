"""Turning a stored file into the plain text that will be indexed."""

import io

from apps.common import media_types
from apps.ingestion import imaging

MIN_CHARS_PER_PAGE = 80
HEADER_CANDIDATE_ROWS = 5
LONG_CELL = 50


class ExtractionError(Exception):
    """Raised when a file cannot be turned into text."""


def can_extract(media_type):
    """Report whether this build can turn a file of one type into text.

    Takes the media type. This is the single answer to what may be indexed:
    the panel reads it to decide whether to offer the switch at all, and the
    switch refuses on it, so a file nobody can read is never queued only to
    fail an hour later with the owner none the wiser.
    """
    return (
        media_type
        in {media_types.PDF, media_types.WORD, media_types.EXCEL, media_types.EXCEL_LEGACY}
        or media_type.startswith("image/")
        or media_type.startswith("text/")
        or media_type in media_types.TEXT_LIKE
    )


def extract_text(data, media_type):
    """Extract the indexable text of a file.

    Takes the raw bytes and the media type, and picks the cheapest route that
    works: a text layer is read directly, and only a file without one goes
    through recognition, so nothing is spent on files that can be read for
    free. Returns the text. Raises ExtractionError for a type this build
    cannot read.
    """
    if media_type == media_types.PDF:
        return extract_pdf(data)
    if media_type.startswith("image/"):
        return imaging.describe(data, media_type)
    if media_type == media_types.WORD:
        return extract_word(data)
    if media_type in {media_types.EXCEL, media_types.EXCEL_LEGACY}:
        return extract_excel(data, media_type)
    if media_type.startswith("text/") or media_type in media_types.TEXT_LIKE:
        return data.decode("utf-8", errors="replace")
    raise ExtractionError(f"No extractor is available for '{media_type}'")


def extract_pdf(data):
    """Read a PDF, falling back to recognition when it is scanned.

    Takes the raw bytes. A scanned page carries an image and no characters, so
    when the text layer yields almost nothing the whole file goes through
    recognition instead. Returns the text.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if len(text) >= MIN_CHARS_PER_PAGE * max(len(pages), 1):
        return text
    return imaging.describe(data, media_types.PDF)


def extract_word(data):
    """Read the paragraphs and tables of a Word document.

    Takes the raw bytes and returns the text. Table cells are included because
    a contract's figures often live only there.
    """
    from docx import Document as WordDocument

    document = WordDocument(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def extract_excel(data, media_type=media_types.EXCEL):
    """Read every sheet of a workbook as records that explain themselves.

    Takes the raw bytes and the media type, since the old binary format needs
    another reader. Returns the text: one paragraph per row, each cell written
    as "heading: value". A row of bare figures halfway down a sheet says
    nothing on its own, and a chunk cut from the middle of a long sheet would
    otherwise carry values with no way to tell a price from a quantity. Rows
    are paragraphs rather than lines so that a chunk boundary, which prefers
    the end of a paragraph, lands between records and never inside one.
    """
    sheets = []
    for title, rows in read_sheets(data, media_type):
        rows = [row for row in (clean_row(row) for row in rows) if any(row)]
        if not rows:
            continue
        header_at = detect_header(rows)
        headings = unique_headings(rows[header_at])
        records = [
            " | ".join(
                f"{heading}: {cell}" if heading else cell
                for heading, cell in zip(headings, row)
                if cell
            )
            for row in rows[header_at + 1 :]
        ]
        records = [record for record in records if record]
        if records:
            sheets.append(f"## {title}\n\n" + "\n\n".join(records))
    return "\n\n".join(sheets)


def read_sheets(data, media_type):
    """Yield (title, rows) for every sheet, whichever format the file is in.

    The modern format is read with computed values rather than the formulas
    that produced them, since the formula is not what anyone asks about.
    """
    if media_type == media_types.EXCEL_LEGACY:
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            yield sheet.name, (sheet.row_values(index) for index in range(sheet.nrows))
        return
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for sheet in book.worksheets:
        yield sheet.title, sheet.iter_rows(values_only=True)


def clean_row(row):
    """Turn the raw cells of one row into the strings worth keeping."""
    return [clean_cell(value) for value in row]


def clean_cell(value):
    """Return a cell as text, or nothing for a cell that carries nothing.

    A float that is a whole number is written without its decimal point, since
    a spreadsheet stores the number 270 that way and nobody searches for 270.0.
    A formula that was never computed is worth nothing to a reader.
    """
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:
            return ""
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    if text.startswith("="):
        return ""
    return text


def detect_header(rows):
    """Find the row that names the columns.

    Takes the cleaned rows. Returns the index of the best candidate among the
    first few: the one with the most short cells that are words rather than
    numbers, which is what a row of headings looks like and a row of data
    usually does not. A title sitting above the real headings, which is
    common, scores below them because it is one cell; a row of figures scores
    below them because figures are not names. Ties go to the earlier row.
    """
    best, best_score = 0, -1
    for index, row in enumerate(rows[:HEADER_CANDIDATE_ROWS]):
        cells = [cell for cell in row if cell]
        score = sum(1 for cell in cells if not is_number(cell))
        if cells and all(len(cell) < LONG_CELL for cell in cells):
            score += 2
        if score > best_score:
            best, best_score = index, score
    return best


def is_number(text):
    """Report whether a cell holds a figure rather than a name."""
    try:
        float(text.replace(",", "."))
    except ValueError:
        return False
    return True


def unique_headings(row):
    """Make every heading distinct, so two columns called Total stay apart."""
    seen = {}
    headings = []
    for index, cell in enumerate(row):
        name = cell or f"column_{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        headings.append(name)
    return headings
